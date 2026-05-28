import openpyxl
import os

def read_xlsx(path, out_file):
    out_file.write(f"=== Reading {path} ===\n")
    if not os.path.exists(path):
        out_file.write("File does not exist\n")
        return
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        out_file.write(f"Sheets: {wb.sheetnames}\n")
        for name in wb.sheetnames:
            sheet = wb[name]
            out_file.write(f"\nSheet: {name}\n")
            # Print first 50 rows
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= 50:
                    break
                # Filter out all-None rows
                if all(v is None for v in row):
                    continue
                row_str = ", ".join([str(v) if v is not None else "" for v in row])
                out_file.write(f"Row {i}: {row_str}\n")
    except Exception as e:
        out_file.write(f"Error: {e}\n")

with open("excel_info.txt", "w", encoding="utf-8") as f:
    read_xlsx("data/消費分類關鍵字對照表.xlsx", f)
    read_xlsx("vba/信用卡明細彙整表_解答.xlsm", f)

print("Done! Check excel_info.txt")
